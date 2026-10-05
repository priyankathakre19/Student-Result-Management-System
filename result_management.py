import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
import os


# ============================================================
# EXCEL FILE SETTINGS
# ============================================================

EXCEL_FILE = "student_results.xlsx"

HEADERS = [
    "Name",
    "Roll No.",
    "Class",
    "Subject 1",
    "Subject 2",
    "Subject 3",
    "Subject 4",
    "Subject 5",
    "Total Marks",
    "Percentage",
    "Result"
]


# ============================================================
# CREATE EXCEL FILE
# ============================================================

def create_excel_file():

    if not os.path.exists(EXCEL_FILE):

        workbook = Workbook()

        sheet = workbook.active
        sheet.title = "Student Results"

        sheet.append(HEADERS)

        workbook.save(EXCEL_FILE)


# Create Excel file when program starts
create_excel_file()


# ============================================================
# MAIN WINDOW
# ============================================================

root = tk.Tk()

root.title("Student Result Management System")
root.geometry("850x600")
root.resizable(False, False)


# ============================================================
# TITLE
# ============================================================

title = tk.Label(
    root,
    text="🎓 STUDENT RESULT MANAGEMENT",
    font=("Arial", 24, "bold")
)

title.pack(pady=(40, 10))


subtitle = tk.Label(
    root,
    text="Manage student results using Python and Excel",
    font=("Arial", 13)
)

subtitle.pack(pady=(0, 30))


# ============================================================
# ADD STUDENT FUNCTION
# ============================================================

def add_student():

    add_window = tk.Toplevel(root)

    add_window.title("Add Student")

    add_window.geometry("550x650")

    add_window.resizable(False, False)


    # --------------------------------------------------------
    # Heading
    # --------------------------------------------------------

    heading = tk.Label(
        add_window,
        text="📝 Add Student",
        font=("Arial", 22, "bold")
    )

    heading.pack(pady=20)


    # --------------------------------------------------------
    # Form Frame
    # --------------------------------------------------------

    form_frame = tk.Frame(add_window)

    form_frame.pack(pady=5)


    # --------------------------------------------------------
    # Entry Variables
    # --------------------------------------------------------

    name_entry = tk.Entry(
        form_frame,
        font=("Arial", 12),
        width=25
    )

    roll_entry = tk.Entry(
        form_frame,
        font=("Arial", 12),
        width=25
    )

    class_entry = tk.Entry(
        form_frame,
        font=("Arial", 12),
        width=25
    )

    sub1_entry = tk.Entry(
        form_frame,
        font=("Arial", 12),
        width=25
    )

    sub2_entry = tk.Entry(
        form_frame,
        font=("Arial", 12),
        width=25
    )

    sub3_entry = tk.Entry(
        form_frame,
        font=("Arial", 12),
        width=25
    )

    sub4_entry = tk.Entry(
        form_frame,
        font=("Arial", 12),
        width=25
    )

    sub5_entry = tk.Entry(
        form_frame,
        font=("Arial", 12),
        width=25
    )


    # --------------------------------------------------------
    # Labels and Entries
    # --------------------------------------------------------

    fields = [
        ("Name:", name_entry),
        ("Roll No.:", roll_entry),
        ("Class:", class_entry),
        ("Subject 1 Marks:", sub1_entry),
        ("Subject 2 Marks:", sub2_entry),
        ("Subject 3 Marks:", sub3_entry),
        ("Subject 4 Marks:", sub4_entry),
        ("Subject 5 Marks:", sub5_entry)
    ]


    for row, (label_text, entry) in enumerate(fields):

        label = tk.Label(
            form_frame,
            text=label_text,
            font=("Arial", 12)
        )

        label.grid(
            row=row,
            column=0,
            padx=10,
            pady=8,
            sticky="w"
        )

        entry.grid(
            row=row,
            column=1,
            padx=10,
            pady=8
        )


    # --------------------------------------------------------
    # Clear Form
    # --------------------------------------------------------

    def clear_form():

        name_entry.delete(0, tk.END)
        roll_entry.delete(0, tk.END)
        class_entry.delete(0, tk.END)

        sub1_entry.delete(0, tk.END)
        sub2_entry.delete(0, tk.END)
        sub3_entry.delete(0, tk.END)
        sub4_entry.delete(0, tk.END)
        sub5_entry.delete(0, tk.END)

        name_entry.focus()


    # --------------------------------------------------------
    # SAVE STUDENT
    # --------------------------------------------------------

    def save_student():

        name = name_entry.get().strip()
        roll_no = roll_entry.get().strip()
        student_class = class_entry.get().strip()


        # Check basic information

        if not name or not roll_no or not student_class:

            messagebox.showerror(
                "Missing Information",
                "Please fill in Name, Roll No. and Class.",
                parent=add_window
            )

            return


        # Check Roll No.

        try:

            int(roll_no)

        except ValueError:

            messagebox.showerror(
                "Invalid Roll No.",
                "Roll No. must contain numbers only.",
                parent=add_window
            )

            return


        # Get marks

        try:

            sub1 = float(sub1_entry.get())
            sub2 = float(sub2_entry.get())
            sub3 = float(sub3_entry.get())
            sub4 = float(sub4_entry.get())
            sub5 = float(sub5_entry.get())

        except ValueError:

            messagebox.showerror(
                "Invalid Marks",
                "Please enter valid marks for all five subjects.",
                parent=add_window
            )

            return


        # Store marks

        marks = [
            sub1,
            sub2,
            sub3,
            sub4,
            sub5
        ]


        # Check marks range

        if any(mark < 0 or mark > 100 for mark in marks):

            messagebox.showerror(
                "Invalid Marks",
                "Each subject mark must be between 0 and 100.",
                parent=add_window
            )

            return


        # ----------------------------------------------------
        # Check Duplicate Roll Number
        # ----------------------------------------------------

        workbook = load_workbook(EXCEL_FILE)

        sheet = workbook["Student Results"]


        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            existing_roll = row[1]

            if str(existing_roll) == roll_no:

                workbook.close()

                messagebox.showerror(
                    "Duplicate Roll No.",
                    "This Roll No. already exists.",
                    parent=add_window
                )

                return


        # ----------------------------------------------------
        # Calculate Result
        # ----------------------------------------------------

        total = sum(marks)

        percentage = total / 5


        # Passing criteria:
        # Minimum 40 marks in every subject

        if all(mark >= 40 for mark in marks):

            result = "Pass"

        else:

            result = "Fail"


        # ----------------------------------------------------
        # Save Data
        # ----------------------------------------------------

        sheet.append([
            name,
            int(roll_no),
            student_class,
            sub1,
            sub2,
            sub3,
            sub4,
            sub5,
            total,
            percentage,
            result
        ])


        workbook.save(EXCEL_FILE)

        workbook.close()


        # ----------------------------------------------------
        # Success Message
        # ----------------------------------------------------

        messagebox.showinfo(
            "Student Saved",
            f"Student saved successfully!\n\n"
            f"Name: {name}\n"
            f"Roll No.: {roll_no}\n"
            f"Total Marks: {total}\n"
            f"Percentage: {percentage:.2f}%\n"
            f"Result: {result}",
            parent=add_window
        )


        # Clear form after saving

        clear_form()


    # --------------------------------------------------------
    # Buttons
    # --------------------------------------------------------

    button_frame = tk.Frame(add_window)

    button_frame.pack(pady=25)


    save_button = tk.Button(
        button_frame,
        text="💾 Save",
        font=("Arial", 12, "bold"),
        width=12,
        command=save_student
    )

    save_button.grid(
        row=0,
        column=0,
        padx=10
    )


    clear_button = tk.Button(
        button_frame,
        text="🔄 Clear",
        font=("Arial", 12, "bold"),
        width=12,
        command=clear_form
    )

    clear_button.grid(
        row=0,
        column=1,
        padx=10
    )


# ============================================================
# GET RESULT FUNCTION
# ============================================================

def get_result():

    result_window = tk.Toplevel(root)

    result_window.title("Get Student Result")

    result_window.geometry("700x450")

    result_window.resizable(False, False)


    # --------------------------------------------------------
    # Heading
    # --------------------------------------------------------

    heading = tk.Label(
        result_window,
        text="🔍 Get Student Result",
        font=("Arial", 22, "bold")
    )

    heading.pack(pady=25)


    # --------------------------------------------------------
    # Search Frame
    # --------------------------------------------------------

    search_frame = tk.Frame(result_window)

    search_frame.pack(pady=10)


    tk.Label(
        search_frame,
        text="Enter Roll No.:",
        font=("Arial", 12)
    ).grid(
        row=0,
        column=0,
        padx=10
    )


    roll_entry = tk.Entry(
        search_frame,
        font=("Arial", 12),
        width=20
    )

    roll_entry.grid(
        row=0,
        column=1,
        padx=10
    )


    # --------------------------------------------------------
    # Result Area
    # --------------------------------------------------------

    result_frame = tk.Frame(result_window)

    result_frame.pack(
        fill="both",
        expand=True,
        padx=20,
        pady=20
    )


    # --------------------------------------------------------
    # Search Function
    # --------------------------------------------------------

    def search_result():

        roll_no = roll_entry.get().strip()


        if not roll_no:

            messagebox.showwarning(
                "Input Required",
                "Please enter a Roll No.",
                parent=result_window
            )

            return


        workbook = load_workbook(EXCEL_FILE)

        sheet = workbook["Student Results"]


        found_student = None


        for row in sheet.iter_rows(
            min_row=2,
            values_only=True
        ):

            if str(row[1]) == roll_no:

                found_student = row

                break


        workbook.close()


        # Clear previous result

        for widget in result_frame.winfo_children():

            widget.destroy()


        # ----------------------------------------------------
        # Student Found
        # ----------------------------------------------------

        if found_student:

            name = found_student[0]
            roll = found_student[1]
            student_class = found_student[2]
            total = found_student[8]
            percentage = found_student[9]
            result = found_student[10]


            details = [
                ("Name", name),
                ("Roll No.", roll),
                ("Class", student_class),
                ("Total Marks", total),
                ("Percentage", f"{float(percentage):.2f}%"),
                ("Result", result)
            ]


            for row, (label, value) in enumerate(details):

                tk.Label(
                    result_frame,
                    text=f"{label}:",
                    font=("Arial", 12, "bold")
                ).grid(
                    row=row,
                    column=0,
                    padx=20,
                    pady=8,
                    sticky="w"
                )


                tk.Label(
                    result_frame,
                    text=value,
                    font=("Arial", 12)
                ).grid(
                    row=row,
                    column=1,
                    padx=20,
                    pady=8,
                    sticky="w"
                )


        else:

            tk.Label(
                result_frame,
                text="❌ Student record not found.",
                font=("Arial", 14, "bold")
            ).pack(pady=30)


    # --------------------------------------------------------
    # Search Button
    # --------------------------------------------------------

    search_button = tk.Button(
        search_frame,
        text="🔍 Get Result",
        font=("Arial", 12, "bold"),
        command=search_result
    )

    search_button.grid(
        row=0,
        column=2,
        padx=10
    )


# ============================================================
# SHOW ALL RESULTS FUNCTION
# ============================================================

def show_all_results():

    all_window = tk.Toplevel(root)

    all_window.title("All Student Results")

    all_window.geometry("1100x550")

    all_window.resizable(True, True)


    # --------------------------------------------------------
    # Heading
    # --------------------------------------------------------

    heading = tk.Label(
        all_window,
        text="📋 All Student Results",
        font=("Arial", 22, "bold")
    )

    heading.pack(pady=20)


    # --------------------------------------------------------
    # Treeview Frame
    # --------------------------------------------------------

    table_frame = tk.Frame(all_window)

    table_frame.pack(
        fill="both",
        expand=True,
        padx=20,
        pady=10
    )


    # --------------------------------------------------------
    # Columns
    # --------------------------------------------------------

    columns = (
        "Name",
        "Roll No.",
        "Class",
        "Total Marks",
        "Percentage",
        "Result"
    )


    tree = ttk.Treeview(
        table_frame,
        columns=columns,
        show="headings",
        height=18
    )


    # --------------------------------------------------------
    # Column Headings
    # --------------------------------------------------------

    for column in columns:

        tree.heading(
            column,
            text=column
        )


    # --------------------------------------------------------
    # Column Widths
    # --------------------------------------------------------

    tree.column(
        "Name",
        width=180,
        anchor="center"
    )

    tree.column(
        "Roll No.",
        width=100,
        anchor="center"
    )

    tree.column(
        "Class",
        width=120,
        anchor="center"
    )

    tree.column(
        "Total Marks",
        width=130,
        anchor="center"
    )

    tree.column(
        "Percentage",
        width=130,
        anchor="center"
    )

    tree.column(
        "Result",
        width=120,
        anchor="center"
    )


    # --------------------------------------------------------
    # Scrollbar
    # --------------------------------------------------------

    scrollbar = ttk.Scrollbar(
        table_frame,
        orient="vertical",
        command=tree.yview
    )

    tree.configure(
        yscrollcommand=scrollbar.set
    )


    tree.pack(
        side="left",
        fill="both",
        expand=True
    )


    scrollbar.pack(
        side="right",
        fill="y"
    )


    # --------------------------------------------------------
    # Read Excel
    # --------------------------------------------------------

    workbook = load_workbook(EXCEL_FILE)

    sheet = workbook["Student Results"]


    student_count = 0


    for row in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        if row[0] is not None:

            name = row[0]
            roll_no = row[1]
            student_class = row[2]
            total = row[8]
            percentage = row[9]
            result = row[10]


            tree.insert(
                "",
                "end",
                values=(
                    name,
                    roll_no,
                    student_class,
                    total,
                    f"{float(percentage):.2f}%",
                    result
                )
            )


            student_count += 1


    workbook.close()


    # --------------------------------------------------------
    # Student Count
    # --------------------------------------------------------

    count_label = tk.Label(
        all_window,
        text=f"Total Students: {student_count}",
        font=("Arial", 12, "bold")
    )

    count_label.pack(
        pady=10
    )


# ============================================================
# EXIT FUNCTION
# ============================================================

def exit_program():

    answer = messagebox.askyesno(
        "Exit",
        "Are you sure you want to exit?"
    )

    if answer:

        root.destroy()


# ============================================================
# MAIN MENU BUTTONS
# ============================================================

button_frame = tk.Frame(root)

button_frame.pack()


# ------------------------------------------------------------
# Add Student
# ------------------------------------------------------------

add_button = tk.Button(
    button_frame,
    text="📝 Add Student",
    font=("Arial", 14, "bold"),
    width=22,
    height=2,
    command=add_student
)

add_button.grid(
    row=0,
    column=0,
    padx=15,
    pady=15
)


# ------------------------------------------------------------
# Get Result
# ------------------------------------------------------------

result_button = tk.Button(
    button_frame,
    text="🔍 Get Result",
    font=("Arial", 14, "bold"),
    width=22,
    height=2,
    command=get_result
)

result_button.grid(
    row=0,
    column=1,
    padx=15,
    pady=15
)


# ------------------------------------------------------------
# Show All Results
# ------------------------------------------------------------

show_button = tk.Button(
    button_frame,
    text="📋 Show All Results",
    font=("Arial", 14, "bold"),
    width=22,
    height=2,
    command=show_all_results
)

show_button.grid(
    row=1,
    column=0,
    padx=15,
    pady=15
)


# ------------------------------------------------------------
# Exit
# ------------------------------------------------------------

exit_button = tk.Button(
    button_frame,
    text="🚪 Exit",
    font=("Arial", 14, "bold"),
    width=22,
    height=2,
    command=exit_program
)

exit_button.grid(
    row=1,
    column=1,
    padx=15,
    pady=15
)


# ============================================================
# FOOTER
# ============================================================

footer = tk.Label(
    root,
    text="Python • Tkinter • OpenPyXL • Excel",
    font=("Arial", 10)
)

footer.pack(
    side="bottom",
    pady=20
)


# ============================================================
# START APPLICATION
# ============================================================

root.mainloop()